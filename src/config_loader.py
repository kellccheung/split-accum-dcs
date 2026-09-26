"""Load DCS run configuration from a single-sheet Excel workbook."""

from __future__ import annotations

import math
from dataclasses import dataclass
from pathlib import Path
from typing import List

from openpyxl import load_workbook

REQUIRED_COLUMNS = (
    "dcs_folder",
    "control_folder",
    "split_1_path_1",
    "split_1_path_2",
    "split_2_path_1",
    "split_2_path_2",
    "split_3_path_1",
    "split_3_path_2",
    "timeout_minutes",
)


@dataclass(frozen=True)
class RunConfig:
    """One Excel row describing a DCS automation run."""

    dcs_folder: str
    control_folder: str
    split_1_path_1: str
    split_1_path_2: str
    split_2_path_1: str
    split_2_path_2: str
    split_3_path_1: str
    split_3_path_2: str
    timeout_minutes: float
    row_number: int

    def param_rows(self) -> List[tuple[int, str, str]]:
        """Return (SPLIT, path_1, path_2) rows for param.csv."""
        return [
            (1, self.split_1_path_1, self.split_1_path_2),
            (2, self.split_2_path_1, self.split_2_path_2),
            (3, self.split_3_path_1, self.split_3_path_2),
        ]


def _cell_str(value: object) -> str:
    if value is None:
        return ""
    return str(value).strip()


def _parse_timeout_minutes(raw: str, excel_row_number: int) -> float:
    """Parse a required positive timeout in minutes."""
    try:
        timeout_minutes = float(raw)
    except ValueError as exc:
        raise ValueError(
            f"Row {excel_row_number}: timeout_minutes must be a positive number, got {raw!r}."
        ) from exc
    if not math.isfinite(timeout_minutes) or timeout_minutes <= 0:
        raise ValueError(
            f"Row {excel_row_number}: timeout_minutes must be a positive number, got {raw!r}."
        )
    return timeout_minutes


def load_run_configs(xlsx_path: Path, sheet_name: str = "runs") -> List[RunConfig]:
    """Read one-row-per-run configs from the given Excel sheet."""
    if not xlsx_path.is_file():
        raise FileNotFoundError(f"Config file not found: {xlsx_path}")

    workbook = load_workbook(xlsx_path, read_only=True, data_only=True)
    if sheet_name not in workbook.sheetnames:
        raise ValueError(
            f"Sheet '{sheet_name}' not found in {xlsx_path.name}. "
            f"Available sheets: {', '.join(workbook.sheetnames)}"
        )

    sheet = workbook[sheet_name]
    rows = sheet.iter_rows(values_only=True)
    try:
        header_row = next(rows)
    except StopIteration as exc:
        raise ValueError(f"Sheet '{sheet_name}' is empty.") from exc

    headers = [_cell_str(cell).lower() for cell in header_row]
    missing = [col for col in REQUIRED_COLUMNS if col not in headers]
    if missing:
        raise ValueError(
            f"Missing required columns in sheet '{sheet_name}': {', '.join(missing)}"
        )

    index = {name: headers.index(name) for name in REQUIRED_COLUMNS}
    configs: List[RunConfig] = []

    for excel_row_number, row in enumerate(rows, start=2):
        if row is None or all(cell is None or _cell_str(cell) == "" for cell in row):
            continue

        values = {
            name: _cell_str(row[index[name]] if index[name] < len(row) else None)
            for name in REQUIRED_COLUMNS
        }
        empty_fields = [name for name, value in values.items() if value == ""]
        if empty_fields:
            raise ValueError(
                f"Row {excel_row_number}: empty required field(s): {', '.join(empty_fields)}"
            )

        timeout_minutes = _parse_timeout_minutes(
            values.pop("timeout_minutes"),
            excel_row_number,
        )
        configs.append(
            RunConfig(
                row_number=excel_row_number,
                timeout_minutes=timeout_minutes,
                **values,
            )
        )

    workbook.close()

    if not configs:
        raise ValueError(f"No data rows found in sheet '{sheet_name}'.")

    return configs
